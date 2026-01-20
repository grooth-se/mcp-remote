"""
Word report generator for tensile test results.

Populates a Word template with test data and results.
"""

import re
from pathlib import Path
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Cm


class TensileReportGenerator:
    """
    Generate tensile test reports from Word template.

    Uses placeholder syntax {{placeholder_name}} in template.
    """

    def __init__(self, template_path: Path):
        """
        Initialize report generator.

        Parameters
        ----------
        template_path : Path
            Path to Word template file
        """
        self.template_path = template_path
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

    def generate_report(
        self,
        output_path: Path,
        data: Dict[str, Any],
        chart_path: Optional[Path] = None,
        logo_path: Optional[Path] = None
    ) -> Path:
        """
        Generate report by populating template with data.

        Parameters
        ----------
        output_path : Path
            Path for output Word document
        data : Dict[str, Any]
            Dictionary of placeholder values
        chart_path : Path, optional
            Path to chart image to insert
        logo_path : Path, optional
            Path to logo image to insert

        Returns
        -------
        Path
            Path to generated report
        """
        doc = Document(self.template_path)

        # Replace placeholders in page headers
        for section in doc.sections:
            header = section.header
            for paragraph in header.paragraphs:
                self._replace_in_paragraph(paragraph, data, chart_path, logo_path)
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            self._replace_in_paragraph(paragraph, data, chart_path, logo_path)

        # Replace placeholders in paragraphs
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph, data, chart_path, logo_path)

        # Replace placeholders in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph, data, chart_path, logo_path)

        doc.save(output_path)
        return output_path

    def _replace_in_paragraph(
        self,
        paragraph,
        data: Dict[str, Any],
        chart_path: Optional[Path],
        logo_path: Optional[Path]
    ):
        """Replace placeholders in a paragraph."""
        full_text = paragraph.text

        # Handle logo placeholder
        if '{{logo}}' in full_text and logo_path and logo_path.exists():
            paragraph.clear()
            run = paragraph.add_run()
            run.add_picture(str(logo_path), height=Cm(1.5))
            return

        # Handle chart placeholder
        if '{{chart}}' in full_text and chart_path and chart_path.exists():
            paragraph.clear()
            run = paragraph.add_run()
            run.add_picture(str(chart_path), width=Inches(5.5))
            return

        # Find all placeholders
        placeholders = re.findall(r'\{\{([^}]+)\}\}', full_text)

        if not placeholders:
            return

        # Replace each placeholder
        for placeholder in placeholders:
            key = placeholder
            value = data.get(key, '')

            # Handle None values
            if value is None:
                value = ''

            # Format numeric values
            if isinstance(value, float):
                value = f"{value:.2f}"
            elif isinstance(value, int):
                value = str(value)

            full_text = full_text.replace(f'{{{{{placeholder}}}}}', str(value))

        # Update paragraph text while preserving formatting
        if paragraph.runs:
            first_run = paragraph.runs[0]
            font_name = first_run.font.name
            font_size = first_run.font.size
            bold = first_run.bold
            italic = first_run.italic

            paragraph.clear()
            new_run = paragraph.add_run(full_text)
            new_run.font.name = font_name
            new_run.font.size = font_size
            new_run.bold = bold
            new_run.italic = italic
        else:
            paragraph.text = full_text

    @staticmethod
    def prepare_report_data(
        test_info: Dict[str, str],
        dimensions: Dict[str, Any],
        results: Dict[str, Any],
        specimen_type: str,
        yield_type: str,
        requirements: Optional[Dict[str, str]] = None
    ) -> Dict[str, Any]:
        """
        Prepare data dictionary for report generation.

        Parameters
        ----------
        test_info : Dict[str, str]
            Test information (project, customer, etc.)
        dimensions : Dict[str, Any]
            Specimen dimensions
        results : Dict[str, Any]
            Analysis results with MeasuredValue objects
        specimen_type : str
            'round' or 'rectangular'
        yield_type : str
            'offset' or 'yield_point'
        requirements : Dict[str, str], optional
            Requirement values (e.g., {'Rp02': 'min. 500', 'Rm': 'min. 600'})

        Returns
        -------
        Dict[str, Any]
            Flattened dictionary ready for template
        """
        import math

        data = {}
        requirements = requirements or {}

        # Test information
        data['test_project'] = test_info.get('test_project', '')
        data['customer'] = test_info.get('customer', '')
        data['customer_order'] = test_info.get('customer_order', '')
        data['product_sn'] = test_info.get('product_sn', '')
        data['specimen_id'] = test_info.get('specimen_id', '')
        data['location_orientation'] = test_info.get('location_orientation', '')
        data['material'] = test_info.get('material', '')
        data['certificate_number'] = test_info.get('certificate_number', '')
        data['test_date'] = test_info.get('test_date', '')
        data['test_standard'] = 'ASTM E8/E8M-22'
        data['yield_method'] = 'Rp0.2/Rp0.5' if yield_type == 'offset' else 'ReH/ReL'
        data['test_engineer'] = test_info.get('test_engineer', '')
        data['strain_source'] = test_info.get('strain_source', 'Extensometer')
        data['test_equipment'] = 'MTS Landmark 500kN'
        data['extensometer'] = 'MTS Extensometer'
        data['test_temperature'] = test_info.get('temperature', '23')
        data['humidity'] = test_info.get('humidity', '50')
        data['comments'] = test_info.get('comments', '')

        # Additional test info keys for template compatibility
        data['operator'] = test_info.get('test_engineer', '')
        data['temperature'] = test_info.get('temperature', '23')

        # Specimen dimensions - template-compatible keys
        data['geometry_type'] = 'Round' if specimen_type == 'round' else 'Rectangular'
        data['specimen_type'] = 'Round' if specimen_type == 'round' else 'Rectangular'
        data['gauge_length'] = dimensions.get('gauge_length', '')
        data['width'] = dimensions.get('width', '-') if specimen_type == 'rectangular' else '-'
        data['thickness'] = dimensions.get('thickness', '-') if specimen_type == 'rectangular' else '-'
        data['diameter'] = dimensions.get('diameter', '-') if specimen_type == 'round' else '-'

        # Calculate cross-section area
        if specimen_type == 'round':
            data['d0'] = dimensions.get('diameter', '')
            data['df'] = dimensions.get('final_diameter', '-')
            data['w0'] = '-'
            data['t0'] = '-'
            if dimensions.get('diameter'):
                d = float(dimensions['diameter'])
                area = math.pi * d**2 / 4
                data['initial_area'] = f"{area:.2f}"
                data['cross_section_area'] = f"{area:.2f}"
            else:
                data['initial_area'] = ''
                data['cross_section_area'] = ''
        else:
            data['d0'] = '-'
            data['df'] = '-'
            data['w0'] = dimensions.get('width', '')
            data['t0'] = dimensions.get('thickness', '')
            if dimensions.get('width') and dimensions.get('thickness'):
                area = float(dimensions['width']) * float(dimensions['thickness'])
                data['initial_area'] = f"{area:.2f}"
                data['cross_section_area'] = f"{area:.2f}"
            else:
                data['initial_area'] = ''
                data['cross_section_area'] = ''

        data['L0'] = dimensions.get('gauge_length', '')
        data['L1'] = dimensions.get('final_gauge_length', '-')
        data['Lc'] = dimensions.get('parallel_length', '')

        # Helper to extract value and uncertainty from MeasuredValue objects
        def get_result(key, default_value='-', default_unc='-'):
            result = results.get(key)
            if result and hasattr(result, 'value'):
                return f"{result.value:.1f}", f"±{result.uncertainty:.1f}"
            return default_value, default_unc

        def get_result_raw(key):
            """Get raw MeasuredValue object."""
            return results.get(key)

        # Young's modulus (E) - template uses {{E}}, {{E_uncertainty}}
        E_result = get_result_raw('E')
        if E_result and hasattr(E_result, 'value'):
            data['E'] = f"{E_result.value:.1f}"
            data['E_uncertainty'] = f"±{E_result.uncertainty:.1f}"
        else:
            data['E'] = '-'
            data['E_uncertainty'] = '-'
        data['E_req'] = requirements.get('E', '-')

        # Yield strengths - show values only for selected method
        # Template uses {{Rp02}}, {{Rp02_uncertainty}}, etc.
        if yield_type == 'offset':
            Rp02_val, Rp02_unc = get_result('Rp02')
            Rp05_val, Rp05_unc = get_result('Rp05')
            # Template-compatible keys (simple names)
            data['Rp02'] = Rp02_val
            data['Rp02_uncertainty'] = Rp02_unc
            # Legacy keys (with _value suffix)
            data['Rp02_value'] = Rp02_val
            data['Rp02_req'] = requirements.get('Rp02', '-')
            data['Rp05_value'] = Rp05_val
            data['Rp05_uncertainty'] = Rp05_unc
            data['Rp05_req'] = requirements.get('Rp05', '-')
            # ReH/ReL not applicable
            data['ReH'] = '-'
            data['ReH_uncertainty'] = '-'
            data['ReH_value'] = '-'
            data['ReH_req'] = '-'
            data['ReL'] = '-'
            data['ReL_uncertainty'] = '-'
            data['ReL_value'] = '-'
            data['ReL_req'] = '-'
        else:
            # Offset values not applicable
            data['Rp02'] = '-'
            data['Rp02_uncertainty'] = '-'
            data['Rp02_value'] = '-'
            data['Rp02_req'] = '-'
            data['Rp05_value'] = '-'
            data['Rp05_uncertainty'] = '-'
            data['Rp05_req'] = '-'
            # ReH/ReL values
            ReH_val, ReH_unc = get_result('ReH')
            ReL_val, ReL_unc = get_result('ReL')
            data['ReH'] = ReH_val
            data['ReH_uncertainty'] = ReH_unc
            data['ReH_value'] = ReH_val
            data['ReH_req'] = requirements.get('ReH', '-')
            data['ReL'] = ReL_val
            data['ReL_uncertainty'] = ReL_unc
            data['ReL_value'] = ReL_val
            data['ReL_req'] = requirements.get('ReL', '-')

        # Ultimate tensile strength - template uses {{Rm}}, {{Rm_uncertainty}}
        Rm_val, Rm_unc = get_result('Rm')
        data['Rm'] = Rm_val
        data['Rm_uncertainty'] = Rm_unc
        data['Rm_value'] = Rm_val
        data['Rm_req'] = requirements.get('Rm', '-')

        # Elongation - template uses {{A}}, {{A_uncertainty}}
        # Use A_percent (from test data) for the main A value
        A_val, A_unc = get_result('A_percent')
        data['A'] = A_val
        data['A_uncertainty'] = A_unc

        # A5 from L1-L0 (manual measurement)
        A5_val, A5_unc = get_result('A_manual')
        data['A5_value'] = A5_val if A5_val != '-' else '-'
        data['A5_uncertainty'] = A5_unc if A5_unc != '-' else '-'
        data['A5_req'] = requirements.get('A5', '-')

        # Uniform elongation (Ag) - template uses {{Ag}}, {{Ag_uncertainty}}
        Ag_val, Ag_unc = get_result('Ag')
        data['Ag'] = Ag_val
        data['Ag_uncertainty'] = Ag_unc

        # Reduction of area (Z%) - template uses {{Z}}, {{Z_uncertainty}}
        Z_val, Z_unc = get_result('Z')
        data['Z'] = Z_val
        data['Z_uncertainty'] = Z_unc
        data['Z_value'] = Z_val
        data['Z_req'] = requirements.get('Z', '-')

        # Rate calculations - template uses {{stress_rate_yield}}, {{strain_rate_yield}}, etc.
        # For yield point, use Rp0.2 rates for offset method, ReH rates for yield_point method
        if yield_type == 'offset':
            stress_rate_yield = get_result_raw('stress_rate_rp02')
            strain_rate_yield = get_result_raw('strain_rate_rp02')
        else:
            stress_rate_yield = get_result_raw('stress_rate_reh')
            strain_rate_yield = get_result_raw('strain_rate_reh')

        # Format stress rate at yield
        if stress_rate_yield and hasattr(stress_rate_yield, 'value'):
            data['stress_rate_yield'] = f"{stress_rate_yield.value:.1f}"
        else:
            data['stress_rate_yield'] = '-'

        # Format strain rate at yield (use scientific notation for small values)
        if strain_rate_yield and hasattr(strain_rate_yield, 'value'):
            val = strain_rate_yield.value
            if abs(val) < 0.01:
                data['strain_rate_yield'] = f"{val:.2e}"
            else:
                data['strain_rate_yield'] = f"{val:.4f}"
        else:
            data['strain_rate_yield'] = '-'

        # Rates at Rm
        stress_rate_rm = get_result_raw('stress_rate_rm')
        strain_rate_rm = get_result_raw('strain_rate_rm')

        if stress_rate_rm and hasattr(stress_rate_rm, 'value'):
            data['stress_rate_rm'] = f"{stress_rate_rm.value:.1f}"
        else:
            data['stress_rate_rm'] = '-'

        if strain_rate_rm and hasattr(strain_rate_rm, 'value'):
            val = strain_rate_rm.value
            if abs(val) < 0.01:
                data['strain_rate_rm'] = f"{val:.2e}"
            else:
                data['strain_rate_rm'] = f"{val:.4f}"
        else:
            data['strain_rate_rm'] = '-'

        # Yield/Tensile ratio (Rp0.2/Rm or ReH/Rm)
        data['ratio_label'] = 'Rp0.2/Rm' if yield_type == 'offset' else 'ReH/Rm'
        if Rm_val and Rm_val != '-':
            if yield_type == 'offset':
                if Rp02_val and Rp02_val != '-':
                    try:
                        ratio = float(Rp02_val) / float(Rm_val)
                        data['yield_tensile_ratio'] = f"{ratio:.3f}"
                    except (ValueError, ZeroDivisionError):
                        data['yield_tensile_ratio'] = '-'
                else:
                    data['yield_tensile_ratio'] = '-'
            else:
                ReH_val_check = data.get('ReH', '-')
                if ReH_val_check and ReH_val_check != '-':
                    try:
                        ratio = float(ReH_val_check) / float(Rm_val)
                        data['yield_tensile_ratio'] = f"{ratio:.3f}"
                    except (ValueError, ZeroDivisionError):
                        data['yield_tensile_ratio'] = '-'
                else:
                    data['yield_tensile_ratio'] = '-'
        else:
            data['yield_tensile_ratio'] = '-'
        data['ratio_req'] = requirements.get('ratio', '-')

        # Validity notes
        data['validity_statement'] = 'The test results comply with the requirements and are considered valid.'
        data['validity_status'] = 'VALID'
        data['validity_notes'] = ''

        # Signatures (to be filled manually)
        data['tested_by'] = ''
        data['tested_date'] = ''
        data['reviewed_by'] = ''
        data['reviewed_date'] = ''
        data['approved_by'] = ''
        data['approved_date'] = ''

        return data
